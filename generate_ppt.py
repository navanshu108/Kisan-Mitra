from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Apply a custom color theme (Green/Earth for agriculture)
BRAND_GREEN = RGBColor(34, 197, 94)
BRAND_DARK_GREEN = RGBColor(20, 83, 45)
GRAY = RGBColor(75, 85, 99)

def style_title(shape, color=BRAND_DARK_GREEN, pt=44):
    for p in shape.text_frame.paragraphs:
        p.font.color.rgb = color
        p.font.bold = True
        p.font.size = Pt(pt)
        p.font.name = "Arial"

def style_body(shape, pt=24):
    for p in shape.text_frame.paragraphs:
        p.font.color.rgb = GRAY
        p.font.size = Pt(pt)
        p.font.name = "Arial"

# Slide 1: Title
slide_layout = prs.slide_layouts[0] # Title slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Kisan Mitra AI"
subtitle.text = "Making government scheme discovery simpler for farmers.\n\nOpen Source Hackathon Project"

style_title(title, pt=54)

# Slide 2: The Problem
slide = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "The Problem"
body.text = (
    "• Information Overload: Too many schemes scattered across portals.\n"
    "• Language Barriers: Complex bureaucratic English alienates farmers.\n"
    "• Eligibility Confusion: Hard for a farmer to know what applies to their specific land & crop.\n"
    "• Lack of Trust: Farmers struggle to verify authentic sources vs fake ones."
)
style_title(title)
style_body(body)

# Slide 3: Our Solution
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "The Solution: Kisan Mitra AI"
body.text = (
    "An intelligent, multilingual matching engine for farmers.\n\n"
    "• Profile-Based: Matches based on State, Crop, and Land Area.\n"
    "• Explains the 'Why': AI explains exactly why a farmer is eligible.\n"
    "• Direct Actions: Links straight to verified official application portals.\n"
    "• Voice & Native Language: Speaks out scheme details in Hindi & Marathi."
)
style_title(title)
style_body(body)

# Slide 4: Key Features Demo
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Key Features"
body.text = (
    "• Smart Categories: Filter by Government, Private, Fertilizer, Machinery, etc.\n"
    "• Official Link Verification: Every scheme links to verified Govt domains.\n"
    "• Multimodal Vision: Upload a physical scheme document and let Gemma AI analyze it.\n"
    "• Complete Offline Fallback: The UI falls back to local data if the API goes offline."
)
style_title(title)
style_body(body)

# Slide 5: Technical Architecture
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Technical Architecture"
body.text = (
    "• Frontend: Vanilla JavaScript SPA + Tailwind CSS for extreme speed and accessibility.\n"
    "• Backend: Python FastAPI handling state logic and matching algorithms.\n"
    "• AI Layer: Google Gemma API integrated for text simplification and Multimodal Vision.\n"
    "• Data: Centralized JSON schema enforcing strict categorization and verified URLs."
)
style_title(title)
style_body(body)

# Slide 6: Hackathon Impact
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Hackathon Alignment"
body.text = (
    "• Best Use of Gemma 4: Leverages Vision and Text capabilities natively.\n"
    "• Open Source AI: Fully open, reproducible, and runnable locally.\n"
    "• Agent Tooling: Developed entirely using agentic engineering and rapid iteration.\n"
    "• Social Impact: Directly empowers the agricultural backbone of the economy."
)
style_title(title)
style_body(body)

# Slide 7: Thank You
slide = prs.slides.add_slide(prs.slide_layouts[0])
title, subtitle = slide.shapes.title, slide.placeholders[1]
title.text = "Thank You!"
subtitle.text = "Try the Demo • Inspect the Code • Empower Farmers"
style_title(title, pt=54)

prs.save("Kisan_Mitra_AI_Presentation.pptx")
print("Presentation successfully generated at Kisan_Mitra_AI_Presentation.pptx")
